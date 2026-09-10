"""엑셀 파일 읽기 — 실적(일일자금계획) / 계획(부서별 자금집행계획서)."""
from __future__ import annotations

import hashlib
import re
from dataclasses import dataclass
from typing import Any

import openpyxl

from core.data_store import Txn

DATE_SHEET = re.compile(r"^(\d{2})(\d{2})(?:-(\d+))?$")   # MMDD, MMDD-N


# ── 실적: 일일자금계획.xlsx ───────────────────────────────────────
def list_date_sheets(path) -> list[tuple[str, int, int, int]]:
    """시트명이 MMDD / MMDD-N 형식인 시트를 전부 찾아 날짜순으로 돌려줍니다.
    반환: [(시트명, 월, 일, 보조번호), ...]"""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    found = []
    for name in wb.sheetnames:
        m = DATE_SHEET.match(name.strip())
        if m:
            found.append((name, int(m.group(1)), int(m.group(2)), int(m.group(3) or 0)))
    wb.close()
    return sorted(found, key=lambda x: (x[1], x[2], x[3]))


def find_month_sheets(path, month: int) -> list[str]:
    """해당 월(MM)로 시작하는 시트를 전부 돌려줍니다. 예: 7월 → 0701, 0702, 0715-2 …"""
    return [n for n, mm, _, _ in list_date_sheets(path) if mm == month]


def available_months(path) -> list[int]:
    """파일 안에 들어 있는 월 목록."""
    return sorted({mm for _, mm, _, _ in list_date_sheets(path)})


def find_last_date_sheet(path) -> str | None:
    """가장 늦은 날짜 시트 하나 (잔고 대사용 — V2에서 사용)."""
    sheets = list_date_sheets(path)
    return sheets[-1][0] if sheets else None



def daily_bundle(path, month: int) -> dict:
    """1번 화면(실적 업로드)에 필요한 걸 **파일 한 번만 열어서** 전부 가져옵니다.

    예전에는 시트 목록·시트별 내용·전월 잔고를 각각 따로 읽느라 화면을 누를 때마다
    같은 엑셀을 9번씩 다시 열었습니다. 큰 파일에서 이게 대기 시간의 대부분이었습니다.

    반환:
      months      파일 안에 있는 월 목록
      sheets      해당 월 시트(날짜순)
      prev_sheet  직전 달 마지막 시트 이름 (없으면 None)
      last_sheet  파일 전체에서 가장 늦은 날짜 시트
      rows        {시트명: 행 목록} — 해당 월 시트 + 전월 마지막 시트
    """
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        found = []
        for name in wb.sheetnames:
            m = DATE_SHEET.match(name.strip())
            if m:
                found.append((name, int(m.group(1)), int(m.group(2)), int(m.group(3) or 0)))
        found.sort(key=lambda x: (x[1], x[2], x[3]))

        months = sorted({mm for _, mm, _, _ in found})
        sheets = [n for n, mm, _, _ in found if mm == month]
        before = [x for x in found if x[1] < month] or [x for x in found if x[1] > month]
        prev_sheet = before[-1][0] if before else None

        last_sheet = found[-1][0] if found else None

        rows = {}
        for name in sheets + ([prev_sheet] if prev_sheet else []):
            if name in rows or name not in wb.sheetnames:
                continue
            ws = wb[name]
            rows[name] = [list(r) for r in ws.iter_rows(values_only=True)]
    finally:
        wb.close()
    return {"months": months, "sheets": sheets, "prev_sheet": prev_sheet,
            "last_sheet": last_sheet, "rows": rows}

HEADER_HINTS = {
    "tag":         ["계정과목", "계정", "과목"],
    "counterpart": ["거래처", "업체", "상대처", "거래처명"],
    "memo":        ["내역", "내용", "적요", "사유", "지출내용", "비고"],
    "amount_in":   ["입금", "수입", "입금액", "수입액"],
    "amount_out":  ["출금", "지출", "지급", "출금액", "지급액", "집행금액", "지출액"],
    "amount":      ["금액", "예상금액", "신청금액"],
    "currency":    ["통화", "환종", "화폐"],
    "date":        ["일자", "날짜", "집행일", "지급일"],
}


def _norm(v: Any) -> str:
    return re.sub(r"\s+", "", str(v)).strip() if v is not None else ""


def detect_columns(rows: list[list], scan_rows: int = 15) -> tuple[int, dict[str, int]]:
    """헤더 행을 찾아 열 번호(0부터)를 추정합니다. 반환: (헤더행 index, {필드: 열번호})"""
    best_row, best_map, best_score = -1, {}, 0
    for i, row in enumerate(rows[:scan_rows]):
        found: dict[str, int] = {}
        for j, cell in enumerate(row):
            text = _norm(cell)
            if not text:
                continue
            for field, hints in HEADER_HINTS.items():
                if field in found:
                    continue
                if any(h in text for h in hints):
                    found[field] = j
                    break
        score = len(found)
        if score > best_score:
            best_row, best_map, best_score = i, found, score
    return best_row, best_map


def mapping_signature(mapping: dict[str, int]) -> str:
    """열 배치 지문 — 이전과 같으면 확인창 없이 자동 적용."""
    s = ",".join(f"{k}:{v}" for k, v in sorted(mapping.items()))
    return hashlib.md5(s.encode()).hexdigest()[:12]


def sheet_rows(path, sheet_name: str | None = None,
               max_rows: int | None = None) -> list[list]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet_name] if sheet_name else wb[wb.sheetnames[0]]
    rows = [list(r) for r in ws.iter_rows(values_only=True, max_row=max_rows)]
    wb.close()
    return rows


def _num(v: Any) -> float:
    if v is None:
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = re.sub(r"[^\d.\-]", "", str(v))
    try:
        return float(s) if s not in ("", "-", ".") else 0.0
    except ValueError:
        return 0.0


def rows_to_txns(rows: list[list], mapping: dict[str, int], header_row: int,
                 period: str, kind: str, dept: str = "",
                 default_currency: str = "KRW") -> list[Txn]:
    """감지된 열 매핑으로 행들을 거래 객체로 변환합니다."""
    out: list[Txn] = []
    for row in rows[header_row + 1:]:
        def g(field: str) -> Any:
            idx = mapping.get(field)
            return row[idx] if idx is not None and idx < len(row) else None

        a_in = _num(g("amount_in"))
        a_out = _num(g("amount_out"))
        if a_in == 0 and a_out == 0:
            # 입/출금 열이 나뉘어 있지 않은 파일 → '금액' 열 사용(방향은 나중에 확인)
            a_out = _num(g("amount"))
        if a_in == 0 and a_out == 0:
            continue

        memo = str(g("memo") or "").strip()
        counterpart = str(g("counterpart") or "").strip()
        if not memo and not counterpart:
            continue

        cur = str(g("currency") or default_currency).strip().upper() or default_currency
        if cur not in ("KRW", "USD", "JPY", "EUR", "CNY"):
            cur = default_currency

        out.append(Txn(
            period=period, kind=kind, dept=dept,
            txn_date=str(g("date") or "").strip(),
            counterpart=counterpart, memo=memo,
            tag=str(g("tag") or "").strip(),
            currency=cur, amount_in=a_in, amount_out=a_out,
            source="upload",
        ))
    return out


# ── 중복 후보 찾기 ────────────────────────────────────────────────
def find_duplicates(txns: list[Txn]) -> list[list[int]]:
    """업체명·적요·금액·방향이 모두 같은 행 2건 이상 → 그룹으로 묶어 반환(인덱스)."""
    buckets: dict[tuple, list[int]] = {}
    for i, t in enumerate(txns):
        key = (t.counterpart, t.memo, round(t.amount, 2), t.direction)
        buckets.setdefault(key, []).append(i)
    return [idx for idx in buckets.values() if len(idx) > 1]


# ── 일일자금계획 전용 파서 ────────────────────────────────────────
# 실제 파일 구조 (2026-08 확인):
#   B=통화구분  C=거래처명  D=내 역  F=수입  G=지출  H=계좌  I=비고
#   한 시트에 표가 둘: "1. 일일자금수지 실적" / "2. 일일자금수지 계획"
#   통화구분 칸에 "1)원화" · "2) 외화(USD)" 가 구간 표시로 들어감
#   각 구간은 "소계" 행에서 끝남
SECTION_ACTUAL = "실적"
SECTION_PLAN = "계획"

_HDR_MUST = ["거래처", "수입", "지출"]


def _find_daily_headers(rows: list[list]) -> list[tuple[int, str, dict[str, int]]]:
    """헤더 행을 전부 찾아 (행번호, 구간종류, 열매핑) 목록으로 돌려줍니다."""
    out = []
    for i, row in enumerate(rows):
        texts = {j: _norm(c) for j, c in enumerate(row) if _norm(c)}
        joined = "".join(texts.values())
        if not all(k in joined for k in _HDR_MUST):
            continue
        mapping: dict[str, int] = {}
        for j, t in texts.items():
            if "통화" in t or "구분" in t:
                mapping.setdefault("currency_section", j)
            elif "거래처" in t:
                mapping.setdefault("counterpart", j)
            elif "내역" in t or "적요" in t or "내용" in t:
                mapping.setdefault("memo", j)
            elif t == "수입" or "입금" in t:
                mapping.setdefault("amount_in", j)
            elif t == "지출" or "출금" in t:
                mapping.setdefault("amount_out", j)
            elif "계좌" in t:
                mapping.setdefault("account", j)
            elif "비고" in t:
                mapping.setdefault("note", j)
            elif "계정" in t or "과목" in t:
                mapping.setdefault("tag", j)
        if "amount_in" not in mapping or "amount_out" not in mapping:
            continue
        # 계정과목 열이 헤더로 안 잡히면(숨김 열이라 제목이 비어 있는 경우 등)
        # 아래 데이터를 훑어서 '아는 계정과목 이름'이 많이 나오는 열을 찾습니다.
        if "tag" not in mapping:
            guess = _guess_tag_column(rows, i, mapping)
            if guess is not None:
                mapping["tag"] = guess
        # 이 표가 실적인지 계획인지 — 위쪽 5행에서 찾음
        kind = SECTION_ACTUAL
        for back in range(max(0, i - 5), i):
            line = "".join(_norm(c) for c in rows[back] if c is not None)
            if "일일자금수지계획" in line or ("계획" in line and "자금수지" in line):
                kind = SECTION_PLAN
            elif "일일자금수지실적" in line or ("실적" in line and "자금수지" in line):
                kind = SECTION_ACTUAL
        out.append((i, kind, mapping))
    return out


def _guess_tag_column(rows: list[list], header_row: int, mapping: dict[str, int],
                      scan: int = 40) -> int | None:
    """헤더에 제목이 없는 계정과목 열을 찾아냅니다(A열이 숨겨져 있는 경우 등).
    태그 사전에 등록된 계정과목 이름이 가장 많이 나오는 열을 고릅니다."""
    from core.engine import ALL_ACCOUNTS
    known = {_norm(a) for a in ALL_ACCOUNTS}
    used = set(mapping.values())
    hits: dict[int, int] = {}
    for row in rows[header_row + 1: header_row + 1 + scan]:
        for j, cell in enumerate(row):
            if j in used:
                continue
            t = _norm(cell)
            if t and t in known:
                hits[j] = hits.get(j, 0) + 1
    if not hits:
        return None
    best = max(hits, key=lambda j: hits[j])
    return best if hits[best] >= 2 else None


def _currency_of(text: str, current: str) -> str:
    t = _norm(text)
    if not t:
        return current
    if "외화" in t or "USD" in t.upper() or "$" in t:
        return "USD"
    if "원화" in t or "KRW" in t.upper():
        return "KRW"
    return current


def _has_currency_marker(rows: list[list], start: int, look: int = 4) -> bool:
    """다음 몇 줄 안에 '2) 외화(USD)' 같은 통화 구간 표시가 있는지."""
    for row in rows[start:start + look]:
        t = "".join(_norm(c) for c in row if c is not None)
        if "외화" in t or "원화" in t or "USD" in t.upper():
            return True
    return False


def parse_daily_sheet(rows: list[list], period: str, sheet_name: str,
                      want: str = SECTION_ACTUAL) -> list[Txn]:
    """일일자금계획 시트 한 장에서 거래를 뽑아냅니다.

    want='실적' 이면 실적 표만, '계획' 이면 계획 표만 읽습니다.
    표는 마지막 '소계' 행에서 끝납니다 — 그 아래 자금대체·계좌잔액 표는 읽지 않습니다.
    """
    headers = _find_daily_headers(rows)
    if not headers:
        return []
    out: list[Txn] = []
    for idx, (hrow, kind, mapping) in enumerate(headers):
        if kind != want:
            continue
        end = headers[idx + 1][0] if idx + 1 < len(headers) else len(rows)
        currency = "KRW"
        for r in range(hrow + 1, end):
            row = rows[r]

            def g(field):
                j = mapping.get(field)
                return row[j] if j is not None and j < len(row) else None

            line = "".join(_norm(c) for c in row if c is not None)
            if not line:
                continue
            currency = _currency_of(str(g("currency_section") or ""), currency)

            # 소계 행 = 구간의 끝. 바로 뒤에 '외화' 같은 통화 표시가 있으면 다음 구간이
            # 이어지는 것이므로 계속 읽고, 없으면 이 표가 끝난 것으로 보고 멈춥니다.
            if "소계" in line or "합계" in line:
                if _has_currency_marker(rows, r + 1):
                    continue
                break

            a_in, a_out = _num(g("amount_in")), _num(g("amount_out"))
            if a_in == 0 and a_out == 0:
                continue
            memo = str(g("memo") or "").strip()
            counterpart = str(g("counterpart") or "").strip()
            if not memo and not counterpart:
                continue

            out.append(Txn(
                period=period,
                kind="actual" if want == SECTION_ACTUAL else "plan",
                txn_date=sheet_name,
                counterpart=counterpart,
                memo=memo,
                tag=str(g("tag") or "").strip(),
                currency=currency,
                amount_in=a_in, amount_out=a_out,
                source="upload",
            ))
    return out


def daily_sheet_summary(rows: list[list]) -> list[dict]:
    """시트 안에 어떤 표가 몇 개 있는지 — 화면에 보여주기 위한 정보."""
    return [{"헤더행": h + 1, "구간": k,
             "열": ", ".join(f"{a}={b}" for a, b in sorted(m.items()))}
            for h, k, m in _find_daily_headers(rows)]


def daily_sheet_subtotals(rows: list[list]) -> list[dict]:
    """시트에 **원본이 직접 적어 둔 '소계' 금액**을 그대로 뽑아옵니다.

    앱이 읽은 합계와 이 값을 나란히 놓으면, 우리가 빠뜨린 줄이 있는지
    사람이 눈으로 바로 확인할 수 있습니다(1번 화면 '원본 소계 대조').
    반환: [{'구간': '실적'|'계획', '통화': 'KRW'|'USD', '수입': …, '지출': …}, …]
    """
    headers = _find_daily_headers(rows)
    out: list[dict] = []
    for idx, (hrow, kind, mapping) in enumerate(headers):
        end = headers[idx + 1][0] if idx + 1 < len(headers) else len(rows)
        currency = "KRW"
        for r in range(hrow + 1, end):
            row = rows[r]

            def g(field):
                j = mapping.get(field)
                return row[j] if j is not None and j < len(row) else None

            line = "".join(_norm(c) for c in row if c is not None)
            if not line:
                continue
            currency = _currency_of(str(g("currency_section") or ""), currency)
            if "소계" in line or "합계" in line:
                out.append({"구간": kind, "통화": currency,
                            "수입": _num(g("amount_in")), "지출": _num(g("amount_out"))})
                if _has_currency_marker(rows, r + 1):
                    continue
                break
    return out


# ── 부서 자금집행계획서 전용 파서 ─────────────────────────────────
# 실제 양식 (2026-08 확인):
#   업체명 | 입/출금 구분 | 사유 | 통화 | 금액 | 부가가치세 | 합계 | 비고(예상 지급일 등)
#   → 집계에 쓰는 금액은 **합계**(부가세 포함). '금액'은 공급가라 쓰면 안 됩니다.
PLAN_FIELDS: list[tuple[str, list[str]]] = [
    # (필드, 후보 키워드 — 앞에 있을수록 우선)
    ("direction",     ["입/출금", "입출금", "입출", "수입지출", "입금출금"]),
    ("counterpart",   ["업체명", "거래처명", "업체", "거래처", "상대처", "지급처"]),
    ("memo",          ["사유", "적요", "내역", "내용", "용도", "지출내용"]),
    ("currency",      ["통화", "환종", "화폐"]),
    ("amount_total",  ["합계", "총액", "지급총액", "지급액계"]),
    ("vat",           ["부가가치세", "부가세", "세액"]),
    ("amount_supply", ["금액", "공급가", "공급가액"]),
    ("note",          ["비고"]),
    ("date",          ["지급예정일", "집행일", "일자", "날짜"]),
]


def detect_plan_columns(rows: list[list], scan_rows: int = 20) -> tuple[int, dict[str, int]]:
    """부서 계획서의 헤더 행과 열 위치를 찾습니다. 반환: (헤더행 index, {필드: 열번호})"""
    best_row, best_map, best_score = -1, {}, 0
    for i, row in enumerate(rows[:scan_rows]):
        texts = {j: _norm(c) for j, c in enumerate(row) if _norm(c)}
        if not texts:
            continue
        mapping: dict[str, int] = {}
        taken: set[int] = set()
        for field, keywords in PLAN_FIELDS:
            for kw in keywords:
                hit = None
                for j, t in sorted(texts.items()):
                    if j in taken:
                        continue
                    if _norm(kw) in t:
                        hit = j
                        break
                if hit is not None:
                    mapping[field] = hit
                    taken.add(hit)
                    break
        score = len(mapping)
        if score > best_score:
            best_row, best_map, best_score = i, mapping, score
    return best_row, best_map


def _direction_of(text: str) -> str | None:
    t = _norm(text)
    if not t:
        return None
    if "출금" in t or "지출" in t or "지급" in t:
        return "out"
    if "입금" in t or "수입" in t:
        return "in"
    return None


def parse_plan_rows(rows: list[list], mapping: dict[str, int], header_row: int,
                    period: str, dept: str, default_currency: str = "KRW",
                    default_direction: str = "out") -> list[Txn]:
    """감지된 열 매핑으로 부서 계획 행을 거래로 바꿉니다.
    금액은 '합계'(부가세 포함)를 우선 사용하고, 없으면 '금액'을 씁니다."""
    out: list[Txn] = []
    for row in rows[header_row + 1:]:
        def g(field):
            j = mapping.get(field)
            return row[j] if j is not None and j < len(row) else None

        line = "".join(_norm(c) for c in row if c is not None)
        if not line:
            continue
        # 합계/소계/총계 행은 건너뜁니다 (예: "합 계 (외화 제외)")
        first_texts = [_norm(row[mapping[f]]) for f in ("counterpart", "memo")
                       if mapping.get(f) is not None and mapping[f] < len(row)]
        if any(t.startswith(("합계", "소계", "총계", "누계")) for t in first_texts):
            continue
        if "소계" in line or "총계" in line:
            continue

        amount = _num(g("amount_total"))
        if amount == 0:
            amount = _num(g("amount_supply"))
        if amount == 0:
            continue

        memo = str(g("memo") or "").strip()
        counterpart = str(g("counterpart") or "").strip()
        if not memo and not counterpart:
            continue
        if _norm(memo) in ("사유", "적요") or _norm(counterpart) in ("업체명", "거래처명"):
            continue   # 헤더가 한 번 더 나온 경우

        direction = _direction_of(str(g("direction") or "")) or default_direction
        cur = _norm(g("currency")).upper() or default_currency
        if cur not in ("KRW", "USD", "JPY", "EUR", "CNY"):
            cur = default_currency

        out.append(Txn(
            period=period, kind="plan", dept=dept,
            txn_date=str(g("date") or g("note") or "").strip(),
            counterpart=counterpart, memo=memo, tag="",
            currency=cur,
            amount_in=amount if direction == "in" else 0.0,
            amount_out=amount if direction == "out" else 0.0,
            source="upload",
        ))
    return out


PLAN_FIELD_LABEL = {
    "counterpart": "업체명", "direction": "입/출금 구분", "memo": "사유",
    "currency": "통화", "amount_supply": "금액(공급가)", "vat": "부가가치세",
    "amount_total": "합계 ← 집계에 사용", "note": "비고", "date": "지급예정일",
}


def column_samples(rows: list[list], header_row: int, col: int, n: int = 3) -> list[str]:
    """어느 열의 실제 값 몇 개를 뽑아옵니다 (화면에서 열을 알아보기 쉽게)."""
    out = []
    for row in rows[header_row + 1:]:
        if col >= len(row):
            continue
        v = row[col]
        if v is None or _norm(v) == "":
            continue
        out.append(f"{v:,.0f}" if isinstance(v, (int, float)) else str(v).strip())
        if len(out) >= n:
            break
    return out


def preview_table(rows: list[list], header_row: int, n: int = 6) -> list[list]:
    """헤더 + 데이터 몇 줄 — 화면에 원본 모양 그대로 보여주기 위한 것."""
    start = max(header_row, 0)
    return [list(r) for r in rows[start:start + n + 1]]


# ── 파일명에서 부서 이름 추측 ─────────────────────────────────────
_DEPT_SUFFIX = ("팀", "실", "본부", "센터", "그룹", "파트", "부")
_DEPT_NOISE = ["자금집행계획서", "자금집행계획", "집행계획서", "계획서", "자금계획",
               "최종", "수정", "final", "ver", "v1", "v2", "복사본"]


def guess_dept(filename: str) -> str:
    """'내부회계관리팀_2026년8월_자금집행계획서.xlsx' → '내부회계관리팀'"""
    name = re.sub(r"\.(xlsx|xlsm|xls)$", "", filename, flags=re.I)
    for noise in _DEPT_NOISE:
        name = re.sub(noise, " ", name, flags=re.I)
    name = re.sub(r"20\d{2}[년\-_.]?\s*\d{1,2}[월\-_.]?", " ", name)
    name = re.sub(r"[\[\]()_\-]+", " ", name)
    parts = [p for p in name.split() if p.strip()]
    for p in parts:
        if p.endswith(_DEPT_SUFFIX) and len(p) >= 2:
            return p
    return parts[0] if parts else "부서명입력"


REQUIRED_FIELDS = ["counterpart", "memo", "amount_total"]
OPTIONAL_FIELDS = ["direction", "currency"]


# ── 시트 고르기 ───────────────────────────────────────────────────
def sheet_names(path, visible_only: bool = False) -> list[str]:
    """시트 이름 목록. visible_only=True 면 숨김 시트는 뺍니다."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    names = [ws.title for ws in wb.worksheets
             if not visible_only or ws.sheet_state == "visible"]
    wb.close()
    return names


def hidden_sheets(path) -> set[str]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out = {ws.title for ws in wb.worksheets if ws.sheet_state != "visible"}
    wb.close()
    return out


_SHEET_NOISE = ("표지", "안내", "양식", "샘플", "예시", "참고", "guide", "cover", "시트")

_MONTH_EN = {"jan": 1, "feb": 2, "mar": 3, "apr": 4, "may": 5, "jun": 6,
             "jul": 7, "aug": 8, "sep": 9, "oct": 10, "nov": 11, "dec": 12}

# 부서마다 시트 이름 규칙이 달라서 여러 형태를 알아듣게 합니다.
_PAT_YM_FULL = re.compile(r"(20\d{2})\s*[년.\-_/]\s*(\d{1,2})\s*월?")   # 2026년 8월 / 2026-08
_PAT_YM_SHORT = re.compile(r"(?<!\d)(\d{2})\s*[년.\-_/]\s*(\d{1,2})\s*월?")  # 26.08
_PAT_YYYYMM = re.compile(r"(?<!\d)(20\d{2})(0[1-9]|1[0-2])(?!\d)")        # 202608
_PAT_YYMM = re.compile(r"(?<!\d)(\d{2})(0[1-9]|1[0-2])(?!\d)")            # 2608
_PAT_M_ONLY = re.compile(r"(?<!\d)(\d{1,2})\s*월")                        # 8월 / 8월계획


def sheet_period(name: str) -> tuple[int | None, int] | None:
    """시트 이름에서 (연도, 월)을 뽑습니다. 연도를 모르면 (None, 월).

    알아듣는 형태: 2026년 8월 · 2026-08 · 26.08 · 202608 · 2608 · 8월 · 8월계획 · Aug
    """
    t = str(name).strip()
    for pat, base in ((_PAT_YM_FULL, 0), (_PAT_YYYYMM, 0)):
        m = pat.search(t)
        if m:
            y, mo = int(m.group(1)), int(m.group(2))
            if 1 <= mo <= 12:
                return (y, mo)
    for pat in (_PAT_YM_SHORT, _PAT_YYMM):
        m = pat.search(t)
        if m:
            y, mo = 2000 + int(m.group(1)), int(m.group(2))
            if 1 <= mo <= 12 and 2000 <= y <= 2099:
                return (y, mo)
    low = t.lower()
    for en, mo in _MONTH_EN.items():
        if en in low:
            m = _PAT_YM_FULL.search(t)
            return (int(m.group(1)) if m else None, mo)
    m = _PAT_M_ONLY.search(t)
    if m and 1 <= int(m.group(1)) <= 12:
        return (None, int(m.group(1)))
    return None


def sheet_matches(name: str, target: tuple[int, int]) -> bool:
    """시트 이름이 목표 연월과 맞는가. 연도가 안 적힌 이름은 월만 맞으면 인정."""
    got = sheet_period(name)
    if not got:
        return False
    y, mo = got
    return mo == target[1] and (y is None or y == target[0])


def _rows_from_ws(ws, max_rows: int | None = None) -> list[list]:
    return [list(r) for r in ws.iter_rows(values_only=True, max_row=max_rows)]


def _score_ws(ws, name: str, target: tuple[int, int] | None = None) -> float:
    """시트 하나가 '계획표'로 얼마나 그럴듯한지 점수 (이미 열린 워크북 기준)."""
    try:
        rows = _rows_from_ws(ws, 60)
    except Exception:
        return -1
    hr, mp = detect_plan_columns(rows)
    score = 0.0
    for f in REQUIRED_FIELDS:
        if mp.get(f) is not None:
            score += 10
    for f in OPTIONAL_FIELDS:
        if mp.get(f) is not None:
            score += 3
    if hr >= 0:
        data = sum(1 for r in rows[hr + 1:] if any(c not in (None, "") for c in r))
        score += min(data, 30) * 0.3
    if any(k in _norm(name).lower() for k in _SHEET_NOISE):
        score -= 15
    if target:
        if sheet_matches(name, target):
            score += 50
        elif sheet_period(name):
            score -= 5
    return score


def score_sheet(path, name: str, target: tuple[int, int] | None = None) -> float:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        return _score_ws(wb[name], name, target)
    finally:
        wb.close()


def pick_best_sheet(path, target: str | None = None,
                    include_hidden: bool = False) -> tuple[str, dict[str, float]]:
    """계획표가 있을 법한 시트를 골라줍니다. **파일은 딱 한 번만 엽니다.**

    탐색 순서 (빠른 것부터):
      1) 이름이 목표 연월과 맞는 시트 → 그것만 내용 확인
      2) 없으면 → 보이는 시트만 내용 확인
      3) 그래도 없으면 → 전체 시트
    """
    want = None
    if target:
        try:
            y, mo = (int(x) for x in target.split("-"))
            want = (y, mo)
        except Exception:
            want = None

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheets = {ws.title: ws for ws in wb.worksheets}
        visible = [ws.title for ws in wb.worksheets if ws.sheet_state == "visible"]
        if include_hidden or not visible:
            visible = list(sheets)

        if want:
            named = [n for n in visible if sheet_matches(n, want)]
            if named:
                scores = {n: _score_ws(sheets[n], n, want) for n in named}
                return max(scores, key=lambda n: scores[n]), scores

        scores = {n: _score_ws(sheets[n], n, want) for n in visible}
        if not scores:
            return "", {}
        return max(scores, key=lambda n: scores[n]), scores
    finally:
        wb.close()


def workbook_overview(path, target: str | None = None,
                      include_hidden: bool = False) -> dict:
    """파일을 **한 번만 열어서** 화면에 필요한 정보를 모두 모아옵니다.

    반환: {"all": [...], "hidden": {...}, "visible": [...],
           "best": 시트명, "scores": {...}, "peek": {시트명: [업체명 …]}}
    """
    want = None
    if target:
        try:
            y, mo = (int(x) for x in target.split("-"))
            want = (y, mo)
        except Exception:
            want = None

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        all_names = [ws.title for ws in wb.worksheets]
        hidden = {ws.title for ws in wb.worksheets if ws.sheet_state != "visible"}
        visible = [n for n in all_names if n not in hidden]
        pool = all_names if (include_hidden or not visible) else visible

        named = [n for n in pool if want and sheet_matches(n, want)]
        targets = named or pool

        sheets = {ws.title: ws for ws in wb.worksheets}
        scores, peek = {}, {}
        for n in targets:
            ws = sheets[n]
            rows = _rows_from_ws(ws, 60)
            hr, mp = detect_plan_columns(rows)
            sc = 0.0
            for f in REQUIRED_FIELDS:
                if mp.get(f) is not None:
                    sc += 10
            for f in OPTIONAL_FIELDS:
                if mp.get(f) is not None:
                    sc += 3
            if hr >= 0:
                data = sum(1 for r in rows[hr + 1:] if any(c not in (None, "") for c in r))
                sc += min(data, 30) * 0.3
            if any(k in _norm(n).lower() for k in _SHEET_NOISE):
                sc -= 15
            if want:
                sc += 50 if sheet_matches(n, want) else (-5 if sheet_period(n) else 0)
            scores[n] = sc
            col = mp.get("counterpart")
            peek[n] = column_samples(rows, hr, col, 3) if (col is not None and hr >= 0) else []

        best = max(scores, key=lambda n: scores[n]) if scores else ""
        return {"all": all_names, "hidden": hidden, "visible": visible,
                "best": best, "scores": scores, "peek": peek}
    finally:
        wb.close()


def sheet_peek(path, name: str, mapping: dict[str, int] | None = None,
               n: int = 3) -> list[str]:
    """시트 안 업체명 몇 개를 미리 보여줍니다 (사용자가 눈으로 확인하도록)."""
    try:
        rows = sheet_rows(path, name, max_rows=40)
    except Exception:
        return []
    hr, mp = detect_plan_columns(rows) if mapping is None else (0, mapping)
    col = mp.get("counterpart")
    if col is None or hr < 0:
        return []
    return column_samples(rows, hr, col, n)


# ── 일일자금계획 「3. 자금현황」 읽기 ─────────────────────────────
# 실제 구조 (2026-08 확인):
#   B열 "3. 자금현황"
#     B "1) 원화 수시입출 예금"  → ... → D "합계"  · H열 당일잔액
#     B "2) 외화(USD) 수시입출 예금" → ... → D "합계"
#     B "3) 운용자금"            → ... → D "합계"
#   ※ 행 번호는 시트마다 다르므로(126행/127행 등) 반드시 글자로 찾습니다.
CASH_SECTIONS = [("krw", ("1)", "원화")), ("usd", ("2)", "외화")), ("invest", ("3)", "운용자금"))]


def _find_col_by_header(rows: list[list], header_row: int, keyword: str) -> int | None:
    for j, cell in enumerate(rows[header_row]):
        if keyword in _norm(cell):
            return j
    return None


def read_cash_status(rows: list[list]) -> dict:
    """「3. 자금현황」에서 구간별 '당일잔액 합계'를 읽습니다.

    반환: {"krw": 원화합계, "usd": 외화합계, "invest": 운용자금합계,
           "rows": {구간: 합계가 있던 행번호}, "found": True/False}
    """
    out = {"krw": None, "usd": None, "invest": None, "rows": {}, "found": False}

    start = None
    for i, row in enumerate(rows):
        line = "".join(_norm(c) for c in row if c is not None)
        if "자금현황" in line:
            start = i
            break
    if start is None:
        return out
    out["found"] = True

    section = None
    total_col = None
    for i in range(start, len(rows)):
        row = rows[i]
        texts = [_norm(c) for c in row if c is not None]
        line = "".join(texts)
        if not line:
            continue

        # 구간 시작 판정
        matched = None
        for key, keys in CASH_SECTIONS:
            if all(k in line for k in keys):
                matched = key
                break
        if matched:
            section = matched
            total_col = None
            continue

        # 구간 안의 헤더 행에서 '당일잔액' 열 위치 찾기
        if section and total_col is None and "당일잔액" in line:
            total_col = _find_col_by_header(rows, i, "당일잔액")
            continue

        # 합계 행
        if section and total_col is not None and any(t == "합계" for t in texts):
            v = row[total_col] if total_col < len(row) else None
            if isinstance(v, (int, float)):
                out[section] = float(v)
                out["rows"][section] = i + 1
            section, total_col = None, None

        # 다음 큰 표(4. …)를 만나면 종료
        if line.startswith("4.") or "일일자금수지" in line and section is None and out["krw"] is not None:
            break
    return out


def read_cash_status_from(path, sheet: str) -> dict:
    return read_cash_status(sheet_rows(path, sheet))


def previous_month_sheet(path, month: int) -> str | None:
    """대상 월 '직전'의 마지막 일자 시트를 찾습니다 (숨김 시트 포함).
    예: 7월이면 0630-1 → 6월 말 잔고를 여기서 읽습니다."""
    sheets = list_date_sheets(path)          # (이름, 월, 일, 보조번호) 날짜순
    before = [s for s in sheets if s[1] < month] or \
             [s for s in sheets if s[1] > month]   # 연도가 넘어간 경우(1월 → 12월)
    return before[-1][0] if before else None
