"""PPT 산출 — 기존 보고 양식(6장)을 템플릿으로 쓰고 **숫자만 갈아 끼웁니다.**

원본 구조 (2026-08 확인, `카탈리스터_7월_자금운용계획_260716__v3.pptx`)
  1장  표지 — 제목·작성일
  2장  표15 = 1. X월 자금수지 / 표10 = 2. 유동자금 계획
  3장  표9  = 3. 운용자금 계획 / 표13 = 4. X월말 자금(stock)
  4장  연간 자금실적(전년) — 건드리지 않음
  5장  연간 자금실적(당해) — 이번 달 칸만 채움
  6장  유동자금 계획 세부내역

표 안의 글꼴·색·테두리는 그대로 두고 **글자만** 바꿉니다.
"""
from __future__ import annotations

import re
from dataclasses import dataclass
from datetime import date
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Pt

from core.config import OUTPUT_DIR, PPT_TEMPLATE_PATH, ROOT
from core.data_store import clean_text
from core.report import Report

# 파일 이름이 달라도 templates 폴더의 .pptx 를 알아서 찾습니다
# (정해진 이름이 없으면 가장 최근 파일 — core/config.find_template)
PPT_TEMPLATE = PPT_TEMPLATE_PATH


@dataclass
class PptReport:
    path: Path
    warnings: list[str]


def fmt(v, dash_zero: bool = True, signed: bool = False) -> str:
    """숫자를 보고서 표기로. 0은 '-', 음수는 -12.3, 증감은 +12.3"""
    if v is None:
        return "-"
    if isinstance(v, str):
        return v
    if dash_zero and abs(v) < 0.05:
        return "-"
    s = f"{abs(v):,.1f}"
    if v < 0:
        return f"-{s}"
    return f"+{s}" if signed else s


def _set_text(cell, text) -> None:
    """셀 글자만 바꾸고 서식(글꼴·크기·색)은 첫 글자 것을 유지합니다."""
    text = clean_text(text)
    tf = cell.text_frame
    para = tf.paragraphs[0]
    if para.runs:
        para.runs[0].text = text
        for extra in para.runs[1:]:
            extra.text = ""
    else:
        para.add_run().text = text
    for p in tf.paragraphs[1:]:
        for r in p.runs:
            r.text = ""


# 비고 칸 글꼴 — 원본이 줄마다 제각각이라 여기서 통일합니다 (사용자 요청 2026-08-25)
NOTE_FONT = "맑은 고딕"
NOTE_SIZE = Pt(8)


def _set_note(cell, text) -> None:
    """비고 칸 — 여러 줄을 **문단으로 나눠** 넣고 글꼴을 8pt 맑은 고딕으로 통일합니다.

    예전에는 줄바꿈이 든 글을 run 하나에 통째로 넣어서, 파워포인트에서 줄이
    안 나뉘고 칸이 늘어나 표 서식이 무너졌습니다.
    """
    text = clean_text(text)
    tf = cell.text_frame
    tf.word_wrap = True
    lines = text.split("\n") if text else [""]

    # 문단 수를 줄 수에 맞춥니다 (남는 문단은 XML에서 제거)
    while len(tf.paragraphs) > len(lines):
        last = tf.paragraphs[-1]._p
        last.getparent().remove(last)
    while len(tf.paragraphs) < len(lines):
        tf.add_paragraph()

    for para, line in zip(tf.paragraphs, lines):
        if para.runs:
            para.runs[0].text = line
            for extra in para.runs[1:]:
                extra.text = ""
        else:
            para.add_run().text = line
        for run in para.runs:
            run.font.name = NOTE_FONT
            run.font.size = NOTE_SIZE


def _tables(slide) -> list:
    return [sh.table for sh in slide.shapes if getattr(sh, "has_table", False)]


def _sub_digits_across_runs(para, pat, repl, group: int = 1) -> bool:
    """문단 전체 글자를 기준으로 찾되, **바뀌는 글자가 든 run만** 고쳐 씁니다.

    파워포인트는 '7월' 한 낱말도 '7' / '월' 처럼 여러 run으로 쪼개 두는 일이 흔합니다
    (실제 원본 확인 결과 슬라이드 안의 월 표기 5곳이 **전부** 쪼개져 있었고,
     그래서 run 단위로만 찾던 예전 코드는 한 곳도 못 바꾸고 있었습니다).
    글꼴·색이 깨지지 않도록 문단을 통째로 합치지 않고, 겹치는 run만 손봅니다.
    """
    runs = list(para.runs)
    if not runs:
        return False
    text = "".join(r.text for r in runs)
    matches = list(pat.finditer(text))
    if not matches:
        return False

    # run 별 [시작, 끝) 위치
    spans, pos = [], 0
    for r in runs:
        spans.append((pos, pos + len(r.text)))
        pos += len(r.text)

    new_texts = [r.text for r in runs]
    for m in reversed(matches):          # 뒤에서부터 고쳐야 위치가 안 밀립니다
        s, e = m.start(group), m.end(group)   # 바뀌는 부분만 골라 고칩니다
        rep = repl(m)
        first = True
        for i, (rs, re_) in enumerate(spans):
            lo, hi = max(s, rs), min(e, re_)
            if lo >= hi:
                continue
            a, b = lo - rs, hi - rs
            cur = new_texts[i]
            new_texts[i] = cur[:a] + (rep if first else "") + cur[b:]
            first = False
    changed = False
    for r, t in zip(runs, new_texts):
        if r.text != t:
            r.text = t
            changed = True
    return changed


def _month_swapper(actual_month: int, plan_month: int, old_actual: int, old_plan: int):
    """옛 달 → 이번 달 바꿔치기 함수. 문단 하나를 받아 고칩니다."""
    # ※ 두 번 이어서 치환하면 '6월→7월→8월'로 연쇄됩니다(실적 7월·계획 8월일 때
    #    제목이 전부 8월이 되던 버그). 한 번의 정규식으로 동시에 바꿉니다.
    table = {str(old_actual): actual_month, str(old_plan): plan_month}
    pat = re.compile(rf"(?<!\d)({old_actual}|{old_plan})월")

    def run(para):
        _sub_digits_across_runs(para, pat, lambda m: str(table[m.group(1)]))
    return run


def _replace_month(shape, actual_month: int, plan_month: int,
                   old_actual: int = 6, old_plan: int = 7) -> None:
    """제목·문장·**표 안**에 박혀 있는 옛 달 숫자를 이번 달로 바꿉니다.

    예전에는 글상자만 고쳐서, 표 머리글의 '7월 계획'·'6월 실적' 이
    지난달 그대로 남아 있었습니다 (2장 유동자금 계획, 3장 운용자금 계획).
    ※ 연간 자금실적 표(4·5장)에는 1~12월 열 이름이 있으므로 **부르면 안 됩니다.**
    """
    swap = _month_swapper(actual_month, plan_month, old_actual, old_plan)
    if getattr(shape, "has_table", False):
        t = shape.table
        for r in range(len(t.rows)):
            for c in range(len(t.columns)):
                for para in t.cell(r, c).text_frame.paragraphs:
                    swap(para)
        return
    if not shape.has_text_frame:
        return
    for para in shape.text_frame.paragraphs:
        swap(para)


_TPL_COVER = re.compile(r"(\d{4})년\s*(\d{1,2})월")
_TPL_AB = re.compile(r"(\d{1,2})\s*월\s*\(\s*([ab])\s*\)")


def detect_template_months(prs) -> tuple[int, int]:
    """템플릿(지난달 보고서)이 **몇 월 실적 / 몇 월 계획**인지 알아냅니다.

    예전에는 6월 실적·7월 계획으로 코드에 박아 뒀습니다. 템플릿을 다른 달 파일로
    바꾸면 월 표기가 안 바뀌므로, 파일에서 직접 읽습니다.
      ① 2장 자금수지 표 머리글 'N월(a)' / 'N월(b)'
      ② 없으면 표지의 '2026년 7월 자금운용계획' → 계획월, 실적월은 그 전달
    """
    slides = list(prs.slides)
    if len(slides) > 1:
        for sh in slides[1].shapes:
            if not getattr(sh, "has_table", False):
                continue
            t = sh.table
            found = {}
            for c in range(len(t.columns)):
                m = _TPL_AB.search(t.cell(0, c).text)
                if m:
                    found[m.group(2)] = int(m.group(1))
            if "a" in found and "b" in found:
                return found["a"], found["b"]
    if slides:
        for sh in slides[0].shapes:
            if not sh.has_text_frame:
                continue
            m = _TPL_COVER.search(sh.text_frame.text)
            if m:
                p = int(m.group(2))
                return (12 if p == 1 else p - 1), p
    return 6, 7


# ── 각 장 맨 위 '요약 한 줄' 채우기 ──────────────────────────────
# 예) "8월 자금수지(flow) : 2,420.5백만 (기초 2,462.5백만 + 자금수지차 -42.0백만
#      = 기말 2,420.5백만)"
# 표가 아니라 자유 문장이지만, 숫자가 전부 '…백만' 형태라 그 숫자만 골라
# 갈아 끼울 수 있습니다. (글자 색·굵기가 run마다 달라서 문단을 통째로
#  덮어쓰지 않고, 숫자가 든 run만 고칩니다.)
_HEADLINE_NUM = re.compile(r"(-?\d[\d,]*(?:\.\d+)?)\s*백만")
_BLUE = RGBColor(0x00, 0x00, 0xFF)     # 음수
_RED = RGBColor(0xD9, 0x2B, 0x2B)      # 증가


def _fill_headline(slide, values, keyword: str, warnings: list, where: str) -> None:
    """values = [(숫자, 부호표시여부), ...] — 문장에 나오는 '…백만' 순서 그대로."""
    target = None
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        for para in sh.text_frame.paragraphs:
            txt = "".join(r.text for r in para.runs)
            if keyword in txt and len(_HEADLINE_NUM.findall(txt)) == len(values):
                target = (para, txt)
                break
        if target:
            break
    if not target:
        warnings.append(f"{where} 맨 위 요약 문장을 못 찾았거나 숫자 개수가 달라 "
                        "그대로 뒀습니다 — 직접 확인해주세요.")
        return

    para, txt = target
    spots = [m.start(1) for m in _HEADLINE_NUM.finditer(txt)]
    ends = [m.end(1) for m in _HEADLINE_NUM.finditer(txt)]
    new_text = {p: fmt(v, dash_zero=False, signed=sg)
                for p, (v, sg) in zip(spots, values)}

    # 색을 다시 칠할 run 찾기 (부호가 있는 칸만). 치환 전에 위치를 잡아둡니다.
    # 숫자 바로 뒤의 '백만'도 원본에서 같은 색이라 함께 칠합니다.
    runs = list(para.runs)
    bounds, pos = [], 0
    for run in runs:
        bounds.append((pos, pos + len(run.text)))
        pos += len(run.text)
    recolor = []
    for i, (a, b) in enumerate(zip(spots, ends)):
        if not values[i][1]:
            continue
        last = None
        for j, (r_start, r_end) in enumerate(bounds):
            if r_start < b and r_end > a:
                recolor.append((runs[j], values[i][0]))
                last = j
        if last is not None and last + 1 < len(runs) \
                and runs[last + 1].text.startswith("백만"):
            recolor.append((runs[last + 1], values[i][0]))

    _sub_digits_across_runs(para, _HEADLINE_NUM,
                            lambda m: new_text[m.start(1)], group=1)

    for run, val in recolor:
        try:
            if run.font.color is not None and run.font.color.type is not None:
                run.font.color.rgb = _BLUE if val < 0 else _RED
        except Exception:
            pass


def build_ppt(rep: Report, out_name: str | None = None,
              title_date: date | None = None,
              template: Path | None = None,
              remarks: dict[str, list[str]] | None = None) -> PptReport:
    """remarks = {"cash_flow": [줄별 비고 …], "liquid_plan": [...], ...}
    미리보기(4번 화면)에서 저장한 비고가 그대로 들어갑니다."""
    warnings: list[str] = []
    tpl = template or PPT_TEMPLATE
    if not Path(tpl).exists():
        raise FileNotFoundError(
            f"PPT 템플릿을 못 찾았습니다: {tpl}\n"
            "지난달 보고 PPT를 templates/자금운용계획_template.pptx 로 복사해 넣어주세요."
        )
    prs = Presentation(str(tpl))
    rm = remarks or {}
    am, pm = rep.actual_month, rep.plan_month
    d = title_date or date.today()

    slides = list(prs.slides)
    if len(slides) < 3:
        warnings.append("PPT 장수가 예상(6장)과 다릅니다. 채우지 못한 표가 있을 수 있습니다.")

    # 템플릿이 몇 월 파일인지 먼저 읽습니다 (고쳐 쓰기 전에!)
    old_a, old_p = detect_template_months(prs)

    def fix_months(slide):
        """글상자와 **표 안**의 옛 달 표기를 이번 달로 바꿉니다.
        (연간 자금실적 4·5장은 1~12월 열이 있어 부르지 않습니다)"""
        for sh in slide.shapes:
            _replace_month(sh, am, pm, old_a, old_p)

    # ── 1장 표지 ──────────────────────────────────────────────────
    if slides:
        _cover_ym = re.compile(r"(\d{4})년\s*(\d{1,2})월")
        _cover_date = re.compile(r"(\d{4})\.\s*(\d{1,2})\.\s*(\d{1,2})")
        for sh in slides[0].shapes:
            if not sh.has_text_frame:
                continue
            whole = sh.text_frame.text
            for para in sh.text_frame.paragraphs:
                if "자금운용계획" in whole:
                    _sub_digits_across_runs(
                        para, _cover_ym, lambda m: f"{d.year}년 {pm}월", group=0)
                _sub_digits_across_runs(
                    para, _cover_date, lambda m: f"{d.year}. {d.month}. {d.day}", group=0)

    # ── 2장: 표15(자금수지) · 표10(유동자금 계획) ─────────────────
    if len(slides) > 1:
        fix_months(slides[1])
        tabs = _tables(slides[1])
        if len(tabs) >= 1:
            t = tabs[0]
            t.cell(0, 1).text_frame.paragraphs[0].runs and _set_text(t.cell(0, 1), f"{am}월(a)")
            _set_text(t.cell(0, 2), f"{pm}월(b)")
            # A기초=1행, B수입=2, C지출=4, D수지차=6, E조달=7, F기말=8
            # 비고는 원본처럼 **실적월(●)·계획월(○)을 위아래 두 줄로** 나눠 씁니다.
            #   B.수입 → 2행(●) / 3행(○),  C.지출 → 4행(●) / 5행(○)
            # (예전에는 한 칸에 몰아넣어서 3·5행에 지난달 문장이 그대로 남았습니다)
            notes = rm.get("cash_flow", [])
            NOTE_ROWS = ((1,), (2, 3), (4, 5), (6,), (7,), (8,))
            for i, (row_i, data) in enumerate(zip((1, 2, 4, 6, 7, 8), rep.tables["cash_flow"])):
                _set_text(t.cell(row_i, 1), fmt(data[1]))
                _set_text(t.cell(row_i, 2), fmt(data[2]))
                _set_text(t.cell(row_i, 3), fmt(data[3], signed=True))
            for i, rows_for in enumerate(NOTE_ROWS):
                txt = notes[i] if i < len(notes) else ""
                lines = [x for x in (txt or "").split("\n")]
                if len(rows_for) == 2 and len(t.rows) > rows_for[1]:
                    _set_note(t.cell(rows_for[0], 4), lines[0] if lines else "")
                    _set_note(t.cell(rows_for[1], 4), "\n".join(lines[1:]))
                elif len(t.rows) > rows_for[0]:
                    _set_note(t.cell(rows_for[0], 4), txt)
        else:
            warnings.append("2장에서 자금수지 표를 못 찾았습니다.")

        if len(tabs) >= 2:
            t = tabs[1]
            notes = rm.get("liquid_plan", [])
            for i, data in enumerate(rep.tables["liquid_plan"]):
                r = 2 + i
                if r >= len(t.rows):
                    break
                for c, v in ((1, data[1]), (2, data[2]), (3, data[3])):
                    _set_text(t.cell(r, c), fmt(v))
                if i < len(notes) and len(t.columns) > 4:
                    _set_note(t.cell(r, 4), notes[i])
        else:
            warnings.append("2장에서 유동자금 계획 표를 못 찾았습니다.")

    # ── 3장: 표9(운용자금) · 표13(월말 자금) ──────────────────────
    if len(slides) > 2:
        fix_months(slides[2])
        tabs = _tables(slides[2])
        if len(tabs) >= 1:
            t = tabs[0]
            notes = rm.get("fund_plan", [])
            for i, data in enumerate(rep.tables["fund_plan"]):
                r = 2 + i
                if r >= len(t.rows):
                    break
                for c, v in zip((2, 3, 4, 5, 6, 7), data[2:8]):
                    _set_text(t.cell(r, c), fmt(v))
                if i < len(notes) and len(t.columns) > 8:
                    _set_note(t.cell(r, 8), notes[i])
        else:
            warnings.append("3장에서 운용자금 계획 표를 못 찾았습니다.")

        if len(tabs) >= 2:
            t = tabs[1]
            _set_text(t.cell(0, 2), f"{am}월(a)")
            _set_text(t.cell(0, 3), f"{pm}월(b)")
            notes = rm.get("stock", [])
            for i, data in enumerate(rep.tables["stock"]):
                r = 1 + i
                if r >= len(t.rows):
                    break
                _set_text(t.cell(r, 2), fmt(data[2]))
                _set_text(t.cell(r, 3), fmt(data[3]))
                _set_text(t.cell(r, 4), fmt(data[4], signed=True))
                if i < len(notes) and len(t.columns) > 5:
                    _set_note(t.cell(r, 5), notes[i])
        else:
            warnings.append("3장에서 월말 자금(stock) 표를 못 찾았습니다.")

    # ── 5장: 연간 자금실적(당해) — 이번 달 칸만 ───────────────────
    if len(slides) > 4:
        tabs = _tables(slides[4])
        if tabs:
            _fill_annual(tabs[0], rep, warnings)
        else:
            warnings.append("5장에서 연간 자금실적 표를 못 찾았습니다.")

    # ── 6장: 유동자금 계획 세부내역 (36줄) ────────────────────────
    if len(slides) > 5:
        fix_months(slides[5])
        tabs = _tables(slides[5])
        if tabs:
            t = tabs[0]
            _set_text(t.cell(0, 4), f"{am}월(실적)")
            _set_text(t.cell(0, 7), f"{pm}월(계획)")
            notes = rm.get("detail", [])
            rows = rep.tables["detail"]
            if len(t.rows) < 2 + len(rows):
                warnings.append(f"6장 세부내역 표의 줄 수가 예상({2+len(rows)}줄)과 다릅니다 "
                                f"— {len(t.rows)}줄. 일부만 채웠습니다.")
            for i, data in enumerate(rows):
                r = 2 + i
                if r >= len(t.rows):
                    break
                for c, v in zip((4, 5, 6, 7, 8, 9), data[1:7]):
                    _set_text(t.cell(r, c), fmt(v))
                _set_text(t.cell(r, 10), fmt(data[7], signed=True))
                if i < len(notes):
                    _set_note(t.cell(r, 11), notes[i])
        else:
            warnings.append("6장에서 세부내역 표를 못 찾았습니다.")

    # ── 각 장 맨 위 요약 문장의 숫자 채우기 ───────────────────────
    if len(slides) > 1:
        cf = rep.tables["cash_flow"]      # [라벨, a, b, b-a]
        _fill_headline(slides[1],
                       [(cf[5][2], False),   # 기말잔고(계획월)
                        (cf[0][2], False),   # 기초잔고
                        (cf[3][2], True),    # 자금수지차
                        (cf[5][2], False)],  # = 기말잔고
                       "flow", warnings, "2장")
    if len(slides) > 2:
        stk = rep.tables["stock"]         # [라벨, 세부, a, b, b-a] / 2·5·8행 = 환산 계
        _fill_headline(slides[2],
                       [(stk[8][3], False),  # 기말자금
                        (stk[8][4], True),   # 전월 대비
                        (stk[2][3], False),  # 유동자금
                        (stk[2][4], True),   # 유동자금 증감
                        (stk[5][3], False)],  # 운용자금
                       "stock", warnings, "3장")

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    path = OUTPUT_DIR / (out_name or f"자금운용계획_{am}월실적_{pm}월계획.pptx")
    prs.save(path)
    return PptReport(path=path, warnings=warnings)


# ══ 5장 연간 자금실적(당해) ═══════════════════════════════════════
# 표 구조 (원본 확인 2026-08-25) — 26행 × 21열
#   r2 : 열 이름 ' 1월 ' … ' 12월 ' ' 합계 ' ' 월평균 '
#   r3 A.기초잔고 / r4 B.수입(경상·비경상) / r7 C.지출(경상·비경상)
#   경상 아래 = 인건비 · S/W,H/W 매입 · 기타(제세공과금~통신비등)
#   r22 D.자금수지차 / r23 D-1.경상수지차 / r24 E.자금조달 / r25 F.기말잔고
#
# 엑셀 '2026년 자금실적' 시트와 **같은 계산값**을 씁니다. 라벨 글자만 조금 달라서
# (예: 엑셀 '경영, 업무지원, 운영용역비' → PPT '경영, 운영용역비') 아래 표로 이어줍니다.
ANNUAL_PPT_LABEL = {
    "경영, 업무지원, 운영용역비": "경영, 운영용역비",
    "식당, 카페, 매점 등 운영비": "식당, 카페 등 운영비",
    "직원비용, 복지포인트 사용비": "직원비용, 복지포인트",
}
# '기타' 아래로 들어가는 경상 그룹 (인건비·S/W,H/W 를 뺀 나머지)
ANNUAL_TOP_GROUPS = ("인건비", "S/W, H/W 매입")


def _nl(s: str) -> str:
    """라벨 비교용 — 공백·쉼표·마침표를 지웁니다."""
    return re.sub(r"[\s,·.]+", "", s or "")


def _annual_month_cols(table) -> tuple[dict[int, int], int | None, int | None]:
    """열 이름 행에서 {월: 열번호} 와 합계·월평균 열을 찾습니다."""
    months: dict[int, int] = {}
    total_col = avg_col = None
    for r in range(min(4, len(table.rows))):
        for c in range(len(table.columns)):
            txt = table.cell(r, c).text.strip()
            m = re.fullmatch(r"(\d{1,2})\s*월", txt)
            if m:
                months.setdefault(int(m.group(1)), c)
            elif txt == "합계" and total_col is None and months:
                total_col = c
            elif txt in ("월평균", "평균") and avg_col is None and months:
                avg_col = c
    return months, total_col, avg_col


def _annual_rows(table) -> list[str]:
    n = min(4, len(table.columns))
    return [" ".join(table.cell(r, c).text.strip() for c in range(n)).strip()
            for r in range(len(table.rows))]


def _row_after(labels: list[str], start: int, want: str) -> int | None:
    w = _nl(want)
    for r in range(start, len(labels)):
        if _nl(labels[r]) == w:
            return r
    return None


def _decimals_in_row(table, row: int, cols: list[int]) -> int:
    """그 줄이 원래 소수점 몇 자리로 적혀 있었는지 (표기 스타일 유지)."""
    best = 0
    for c in cols:
        txt = table.cell(row, c).text.strip().replace(",", "")
        if "." in txt:
            best = max(best, len(txt.split(".")[-1].rstrip()))
    return min(best, 2)


def _annual_num(v: float, dec: int) -> str:
    """0이면 '-', 아주 작은 값은 '0'(원본 표기 방식 그대로).
    1e-6백만원 = 1원 미만은 계산 오차로 보고 0으로 봅니다."""
    if abs(v) < 1e-6:
        return "-"
    return f"{v:,.{dec}f}"


def _annual_breakdown(result, rep_fx: float = 0.0) -> dict:
    """한 달치 연간표 값 — 엑셀 자금실적과 같은 기준으로 묶습니다."""
    from core.engine import (ANNUAL_GROUP, EXTRAORDINARY_INCOME_ACCOUNTS,
                             MILLION)

    groups = {g: 0.0 for g, _ in ANNUAL_GROUP.values()}
    cls = {g: c for g, c in ANNUAL_GROUP.values()}
    if result is not None:
        for key, val in result.values.items():
            if key in ANNUAL_GROUP:
                groups[ANNUAL_GROUP[key][0]] += val

    # 수입 경상/비경상 (자산매각대금·투자분배금·투자수익만 비경상)
    extra_in = 0.0
    if result is not None:
        for key, txns in result.rows_by_key.items():
            if key.startswith("income."):
                for tx in txns:
                    if (tx.tag or "").strip() in EXTRAORDINARY_INCOME_ACCOUNTS:
                        extra_in += tx.amount / MILLION
    income = result.values.get("income.total", 0.0) if result is not None else 0.0

    ord_groups = {g: v for g, v in groups.items() if cls.get(g) == "경상"}
    etc = {g: v for g, v in ord_groups.items() if g not in ANNUAL_TOP_GROUPS}
    return {
        "income": income,
        "income_ord": income - extra_in,
        "income_ext": extra_in,
        "groups": groups,
        "etc_groups": etc,
        "expense_ord": sum(ord_groups.values()),
        "expense_ext": sum(v for g, v in groups.items() if cls.get(g) == "비경상"),
        "etc_total": sum(etc.values()),
    }


def _fill_annual(table, rep: Report, warnings: list[str]) -> None:
    """연간 자금실적(당해) 표의 **실적월·계획월 두 칸**을 채웁니다.

    예전에는 실적월의 B.수입·C.지출 두 줄만 고쳐서, 나머지가 지난달 값 그대로
    남아 있었습니다 (사용자 확인 2026-08-25).
    """
    months, total_col, avg_col = _annual_month_cols(table)
    labels = _annual_rows(table)
    cf = rep.tables["cash_flow"]     # [라벨, 실적월, 계획월, 증감]

    def find(want, start=0):
        return _row_after(labels, start, want)

    r_open = find("A. 기초잔고")
    r_in = find("B. 수입")
    r_out = find("C. 지출")
    r_diff = find("D. 자금수지차 (=B-C)")
    r_diff_ord = find("D-1. 경상수지차")
    r_fin = find("E. 자금조달(증자)")
    r_close = find("F. 기말잔고(=A+D+E)")
    if r_in is None or r_out is None:
        warnings.append("5장 연간 자금실적에서 'B. 수입'·'C. 지출' 줄을 못 찾았습니다.")
        return
    r_in_ord = _row_after(labels, r_in + 1, "경상")
    r_in_ext = _row_after(labels, r_in + 1, "비경상")
    r_out_ord = _row_after(labels, r_out + 1, "경상")
    r_out_ext = _row_after(labels, (r_out_ord or r_out) + 1, "비경상")

    data = {rep.actual_month: _annual_breakdown(rep.actual),
            rep.plan_month: _annual_breakdown(rep.plan)}
    col_of = {rep.actual_month: 1, rep.plan_month: 2}     # cash_flow 열

    month_cols = [months[m] for m in sorted(months)] if months else []
    missing_grp: set[str] = set()

    for month, d in data.items():
        col = months.get(month)
        if col is None:
            warnings.append(f"5장 연간 자금실적에서 {month}월 칸을 못 찾았습니다.")
            continue
        i = col_of[month]
        plan: list[tuple[int | None, float]] = [
            (r_open, cf[0][i]), (r_in, cf[1][i]),
            (r_in_ord, d["income_ord"]), (r_in_ext, d["income_ext"]),
            (r_out, cf[2][i]), (r_out_ord, d["expense_ord"]),
            (r_out_ext, d["expense_ext"]),
            (r_diff, cf[3][i]),
            (r_diff_ord, d["income_ord"] - d["expense_ord"]),
            (r_fin, cf[4][i]), (r_close, cf[5][i]),
        ]
        # 경상 아래 세부 그룹
        for grp, val in d["groups"].items():
            name = ANNUAL_PPT_LABEL.get(grp, grp)
            row = find(name)
            if row is None:
                if abs(val) > 1e-9:
                    missing_grp.add(grp)
                continue
            plan.append((row, val))
        row_etc = find("기타", r_out)
        if row_etc is not None:
            plan.append((row_etc, d["etc_total"]))

        for row, val in plan:
            if row is None:
                continue
            dec = _decimals_in_row(table, row, month_cols)
            _set_text(table.cell(row, col), _annual_num(val, dec))

    if missing_grp:
        warnings.append("5장 연간 자금실적에 대응 줄이 없는 항목이 있어 '기타'에만 "
                        "합산했습니다: " + ", ".join(sorted(missing_grp)))

    # 합계·월평균 다시 계산 (기초잔고·기말잔고 줄은 원본 그대로 둡니다)
    if total_col is None or not month_cols:
        return
    flow_rows = {r for r in (r_in, r_in_ord, r_in_ext, r_out, r_out_ord, r_out_ext,
                             r_diff, r_diff_ord, r_fin) if r is not None}
    for grp in set(ANNUAL_PPT_LABEL.values()) | {"인건비", "S/W, H/W 매입"}:
        rr = find(grp)
        if rr is not None:
            flow_rows.add(rr)
    for grp in ("제세공과금", "전산관련지출", "복리후생비", "임차료, 관리비",
                "지급수수료", "통신비등"):
        rr = find(grp)
        if rr is not None:
            flow_rows.add(rr)
    if r_out is not None:
        rr = find("기타", r_out)
        if rr is not None:
            flow_rows.add(rr)

    for row in sorted(flow_rows):
        dec = _decimals_in_row(table, row, month_cols)
        s = 0.0
        for c in month_cols:
            txt = table.cell(row, c).text.strip().replace(",", "")
            try:
                s += float(txt)
            except ValueError:
                pass
        _set_text(table.cell(row, total_col), _annual_num(s, dec))
        if avg_col is not None and avg_col != total_col:
            _set_text(table.cell(row, avg_col), _annual_num(s / 12.0, dec))
