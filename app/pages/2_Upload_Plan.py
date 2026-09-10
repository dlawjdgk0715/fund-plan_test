import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parents[2]))

import pandas as pd
import streamlit as st
from openpyxl.utils import get_column_letter

from core.nav import apply_nav
from core.data_store import (add_tag_example, clear_period_dept, dept_counts,
                             init_db, load_column_mapping, load_tags,
                             save_column_mapping, save_transactions)
from core.parsers import (OPTIONAL_FIELDS, PLAN_FIELD_LABEL, REQUIRED_FIELDS,
                          column_samples, detect_plan_columns, find_duplicates,
                          guess_dept, hidden_sheets, mapping_signature,
                          parse_plan_rows, pick_best_sheet, preview_table,
                          sheet_matches, sheet_names, sheet_peek, sheet_period,
                          sheet_rows, workbook_overview)
from core.uploads import KeptFile, forget, keep, recall, remember


# ── 파일을 한 번만 열도록 기억해둡니다 (Streamlit 캐시) ──────────
# 화면을 누를 때마다 코드가 처음부터 다시 실행되는데, 캐시가 없으면
# 그때마다 엑셀을 다시 엽니다. 파일 내용이 바뀌면 자동으로 다시 읽습니다.
# 올린 파일 자체는 core/uploads.py 가 보관합니다 (화면을 옮겨도 안 사라지도록).
@st.cache_data(show_spinner=False, max_entries=30)
def c_overview(path: str, stamp: str, target: str, include_hidden: bool) -> dict:
    """파일을 한 번만 열어서 시트 목록·숨김 여부·추천 시트·미리보기를 모두 가져옵니다."""
    return workbook_overview(path, target=target, include_hidden=include_hidden)


@st.cache_data(show_spinner=False, max_entries=60)
def c_rows(path: str, stamp: str, sheet: str, max_rows: int | None = None) -> list[list]:
    return sheet_rows(path, sheet, max_rows=max_rows)



from core.period import fmt, next_period
from core.tag_classifier import suggest_tag

st.set_page_config(page_title="2. 계획 업로드", layout="wide")
apply_nav(st)
init_db()
st.title("2️⃣ 계획 업로드 — 부서별 자금집행계획서")

period = st.session_state.get("period", "")
if not period:
    st.error("먼저 시작 화면(Home)에서 실적 월을 정해주세요.")
    st.stop()

st.info(f"📅 **{fmt(next_period(period))} 계획** 파일을 올려주세요 (실적 월 {fmt(period)} 기준)")

with st.expander("❓ 여러 부서 파일을 한꺼번에 올리는 방법"):
    st.markdown("""
**파일을 한 번에 여러 개 끌어다 놓으면 됩니다.** 파일 1개 = 부서 1개로 처리합니다.

**부서명은 파일 이름에서 자동으로 뽑습니다.**

| 파일 이름 | 자동 인식되는 부서명 |
|---|---|
| `내부회계관리팀_2026년8월_자금집행계획서.xlsx` | 내부회계관리팀 |
| `2026.08 인프라실 자금집행계획.xlsx` | 인프라실 |
| `전략금융실-계획서(최종).xlsx` | 전략금융실 |

`팀`·`실`·`본부`·`센터`로 끝나는 낱말을 찾아내고, 날짜나 "자금집행계획서" 같은 공통 단어는 걸러냅니다.
**자동 인식이 틀리면 아래 표에서 직접 고치시면 됩니다.**

⚠️ 부서명은 열 배치를 기억하는 열쇠입니다. **매달 같은 이름을 쓰세요.**
이름이 같으면 다음 달부터 열 확인 없이 자동으로 넘어갑니다.

📌 **시트가 여러 개인 파일**은 표지·옛날 데이터 시트를 걸러내고 계획표가 있는 시트를
자동으로 고릅니다. 자동 선택이 틀리면 아래 '시트 선택'에서 직접 바꾸시면 됩니다.
""")

# ── 올린 파일은 화면을 옮겨도 그대로 남습니다 ─────────────────────
# 스트림릿은 다른 화면에 갔다 오면 업로드 칸을 비워버립니다(원래 동작).
KEEP_KEY = "plan_files"

_ups_widget = st.file_uploader("부서 계획 파일 (여러 개 한꺼번에 가능)",
                               type=["xlsx", "xlsm"], accept_multiple_files=True)
if _ups_widget:
    recs = [keep("plan", u.name, u.getbuffer().tobytes()).to_dict() for u in _ups_widget]
    st.session_state[KEEP_KEY] = recs
    remember("plan", recs)             # cmd 창을 껐다 켜도 남도록

ups = KeptFile.list_from(st.session_state.get(KEEP_KEY))
if not ups:                            # 앱을 새로 켠 경우 — 지난번 파일들을 되살립니다
    ups = KeptFile.list_from(recall("plan"))
    if ups:
        st.session_state[KEEP_KEY] = [u.to_dict() for u in ups]
if not ups:
    st.session_state.pop(KEEP_KEY, None)
    st.stop()

if not _ups_widget:
    k1, k2 = st.columns([4, 1])
    k1.success(f"직전에 올리신 **{len(ups)}개 파일**을 그대로 씁니다 "
               f"({', '.join(u.name for u in ups[:3])}"
               f"{' 외 ' + str(len(ups) - 3) + '개' if len(ups) > 3 else ''}). "
               "바꾸시려면 위에 새로 올려주세요.")
    with k2:
        st.write("")
        if st.button("↺ 파일 비우기", use_container_width=True):
            st.session_state.pop(KEEP_KEY, None)
            st.session_state.pop("plan_txns", None)
            forget("plan")
            st.rerun()

# ── 부서명 확인 ───────────────────────────────────────────────────
st.subheader("① 부서명 확인")
st.caption("파일 이름에서 자동으로 뽑았습니다. 틀리면 '부서명' 칸을 눌러 고쳐주세요.")
dept_df = st.data_editor(
    pd.DataFrame([{"파일": u.name, "부서명": guess_dept(u.name), "포함": True} for u in ups]),
    use_container_width=True, hide_index=True, key="dept_editor",
    column_config={"파일": st.column_config.TextColumn(disabled=True),
                   "포함": st.column_config.CheckboxColumn(help="체크 해제하면 이 파일은 건너뜁니다")},
)

files = []
for u, (_, r) in zip(ups, dept_df.iterrows()):
    if r["포함"] and str(r["부서명"]).strip():
        files.append((str(r["부서명"]).strip(), u))
if not files:
    st.warning("포함할 파일이 없습니다.")
    st.stop()

dups_name = [d for d in {f[0] for f in files} if [f[0] for f in files].count(d) > 1]
if dups_name:
    st.error(f"부서명이 겹칩니다: {', '.join(dups_name)} — 서로 다른 이름으로 고쳐주세요.")
    st.stop()

# ── 시트 선택 ─────────────────────────────────────────────────────
paths, stamps = {}, {}
for dept, u in files:
    # 보관된 파일은 이름 자체가 내용 지문이라, 파일과 내용이 항상 짝이 맞습니다.
    paths[dept] = u.path
    stamps[dept] = u.stamp

plan_period = next_period(period)

ov = {}
with st.spinner("파일을 읽는 중..."):
    for dept, pth in paths.items():
        ov[dept] = c_overview(pth, stamps[dept], plan_period, False)

for dept in paths:
    key = f"sheet_{dept}"
    if key not in st.session_state:
        o = ov[dept]
        st.session_state[key] = o["best"] or (o["visible"] or o["all"])[0]

multi = [d for d in paths if len(ov[d]["all"]) > 1]

if multi:
    st.divider()
    st.subheader("② 시트 선택")
    st.caption(f"시트가 여러 개인 파일입니다. **{fmt(plan_period)}** 계획표가 있는 시트를 "
               "자동으로 골랐습니다. 오른쪽 업체명이 맞는지 확인해주세요.")
    _want = tuple(int(x) for x in plan_period.split("-"))

    for dept in multi:
        o = ov[dept]
        n_hid = len(o["hidden"])
        show_hidden = False
        if n_hid:
            show_hidden = st.checkbox(
                f"{dept} — 숨김 시트 {n_hid}개도 후보에 넣기", key=f"sh_{dept}",
                help="과거 자료 보관용 숨김 시트까지 보고 싶을 때만 켜세요.")
            if show_hidden:
                o = c_overview(paths[dept], stamps[dept], plan_period, True)

        names = o["all"] if show_hidden else (o["visible"] or o["all"])
        cur = st.session_state[f"sheet_{dept}"]
        if cur not in names:
            cur = o["best"] or names[0]

        def _fmt_sheet(n, sc=o["scores"], hd=o["hidden"], w=_want):
            if sheet_matches(n, w):
                return f"{n}  ⭐ 이번 계획월"
            if sc.get(n, 0) >= 25:
                return f"{n}  · 계획표로 보임"
            if sheet_period(n):
                return f"{n}  · 다른 달"
            if n in hd:
                return f"{n}  · 숨김"
            return n

        c1, c2 = st.columns([2, 3])
        with c1:
            picked = st.selectbox(
                f"**{dept}** (전체 {len(o['all'])}개 · 보이는 시트 {len(o['visible'])}개)",
                names, index=names.index(cur), format_func=_fmt_sheet,
                key=f"sheetsel_{dept}")
        with c2:
            st.write("")
            peek = o["peek"].get(picked)
            if peek is None:
                st.caption("업체명 미리보기: (선택 후 아래 표에서 확인)")
            elif peek:
                st.caption("이 시트 업체명: " + " · ".join(peek))
            else:
                st.caption("⚠️ 이 시트에서 업체명을 못 찾았습니다 — 다른 시트일 수 있어요.")

        if picked != st.session_state[f"sheet_{dept}"]:
            st.session_state[f"sheet_{dept}"] = picked
            st.session_state.pop(f"colmap_{dept}", None)
            st.rerun()

# ── 열 감지 ───────────────────────────────────────────────────────
st.divider()
st.subheader("③ 이렇게 읽었습니다" if multi else "② 이렇게 읽었습니다")

info = {}
for dept, u in files:
    tmp = paths[dept]
    rows = c_rows(tmp, stamps[dept], st.session_state[f"sheet_{dept}"])
    det_row, det_map = detect_plan_columns(rows)
    saved = load_column_mapping(dept)
    reused = bool(saved and saved[1] == mapping_signature(det_map))
    if reused:
        det_map = saved[0]
    key = f"colmap_{dept}"
    if key not in st.session_state:
        st.session_state[key] = (det_row, dict(det_map))
    info[dept] = {"rows": rows, "reused": reused,
                  "sheet": st.session_state[f"sheet_{dept}"]}

FIELD_ORDER = [("counterpart", "업체명"), ("memo", "사유"),
               ("amount_total", "💰 합계(집계금액)"), ("direction", "입/출금"),
               ("currency", "통화")]

table = []
for dept, _ in files:
    rows = info[dept]["rows"]
    hr, mp = st.session_state[f"colmap_{dept}"]
    row = {"부서": dept, "시트": info[dept]["sheet"]}
    for f, label in FIELD_ORDER:
        col = mp.get(f)
        if col is None:
            row[label] = "❌ 못 찾음"
        else:
            head = ""
            if 0 <= hr < len(rows) and col < len(rows[hr]):
                head = str(rows[hr][col] or "").strip()
            ex = " · ".join(column_samples(rows, hr, col, 2))
            row[label] = f"✅ {get_column_letter(col+1)}열 「{head}」 → {ex}"
    missing = [l for f, l in FIELD_ORDER
               if f in REQUIRED_FIELDS and mp.get(f) is None]
    row["상태"] = "❌ 확인 필요" if missing else ("♻️ 지난달과 동일" if info[dept]["reused"] else "✅ 정상")
    table.append(row)

tdf = pd.DataFrame(table)


def _color(v):
    v = str(v)
    if v.startswith("❌"):
        return "background-color:#ffe3e3; color:#b00020; font-weight:600"
    if v.startswith("♻️"):
        return "background-color:#e7f5ff; color:#1864ab"
    if v.startswith("✅ 정상"):
        return "background-color:#e6fcf5; color:#0b7285"
    return ""


st.dataframe(tdf.style.map(_color), use_container_width=True, hide_index=True)

bad = [r["부서"] for r in table if r["상태"].startswith("❌")]
if bad:
    st.error(f"🔴 **{', '.join(bad)}** — 꼭 필요한 열(업체명·사유·합계)을 못 찾았습니다. "
             "아래 '고칠 게 있어요'를 눌러 지정해주세요.")

# ── 고치기 버튼 ───────────────────────────────────────────────────
if "fix_open" not in st.session_state:
    st.session_state.fix_open = bool(bad)
label = "🔽 열 지정 화면 닫기" if st.session_state.fix_open else "✏️  고칠 게 있어요 — 열 직접 지정하기"
if st.button(label, type="primary" if not st.session_state.fix_open else "secondary",
             use_container_width=True):
    st.session_state.fix_open = not st.session_state.fix_open
    st.rerun()

ROLES = ["counterpart", "direction", "memo", "currency",
         "amount_total", "amount_supply", "vat", "note", "date"]
ROLE_NONE = "— 안 씀 —"

if st.session_state.fix_open:
    st.divider()
    which = st.selectbox("어느 부서를 고칠까요?", [d for d, _ in files],
                         index=([d for d, _ in files].index(bad[0]) if bad else 0))
    rows = info[which]["rows"]
    hr, mp = st.session_state[f"colmap_{which}"]
    hr = st.number_input("제목(헤더)이 있는 행", 1, min(60, len(rows)), hr + 1,
                         help="업체명·사유·합계 같은 제목이 적힌 줄") - 1

    pv = preview_table(rows, hr, 5)
    total_cols = max((len(r) for r in pv), default=0)
    MAX_SHOW = 12
    if total_cols > MAX_SHOW:
        st.caption(f"이 파일은 열이 {total_cols}개입니다. 한 번에 {MAX_SHOW}개씩 보여드립니다.")
        blocks = (total_cols + MAX_SHOW - 1) // MAX_SHOW
        blk = st.number_input(
            f"열 묶음 ({blocks}개 중)", 1, blocks, 1, key=f"blk_{which}",
            help="찾는 열이 안 보이면 다음 묶음으로 넘겨보세요.") - 1
        lo, hi = blk * MAX_SHOW, min((blk + 1) * MAX_SHOW, total_cols)
    else:
        lo, hi = 0, total_cols
    st.caption(f"지금 보는 열: {get_column_letter(lo+1)} ~ {get_column_letter(hi)}열")

    # 이미 지정된 역할은 화면에 안 보이는 열이라도 그대로 유지합니다
    newmap: dict[str, int] = {r: c for r, c in mp.items() if not (lo <= c < hi)}
    labels = {ROLE_NONE: ROLE_NONE, **{r: PLAN_FIELD_LABEL.get(r, r) for r in ROLES}}
    for start in range(lo, hi, 4):
        cols = st.columns(min(4, hi - start))
        for k, col in enumerate(range(start, min(start + 4, hi))):
            with cols[k]:
                head = str(rows[hr][col]).strip() if hr < len(rows) and col < len(rows[hr]) \
                    and rows[hr][col] is not None else "(제목없음)"
                st.markdown(f"**{get_column_letter(col+1)}열** · {head}")
                st.caption("\n\n".join(f"· {s}" for s in column_samples(rows, hr, col, 3))
                           or "· (빈 열)")
                cur = next((r for r, v in mp.items() if v == col), ROLE_NONE)
                opts = [ROLE_NONE] + ROLES
                sel = st.selectbox("역할", opts, index=opts.index(cur) if cur in opts else 0,
                                   format_func=lambda r: labels[r],
                                   key=f"role_{which}_{col}", label_visibility="collapsed")
                if sel != ROLE_NONE:
                    newmap[sel] = col
    st.session_state[f"colmap_{which}"] = (hr, newmap)

    st.markdown("**원본 미리보기**")
    ncols = hi - lo
    body = [[(r[c] if c < len(r) else None) for c in range(lo, hi)] for r in pv[1:]]
    headers = []
    for c in range(lo, hi):
        title = str(pv[0][c] or "")[:12] if c < len(pv[0]) and pv[0][c] is not None else ""
        headers.append(f"{get_column_letter(c+1)} · {title}" if title
                       else f"{get_column_letter(c+1)}열")
    st.dataframe(pd.DataFrame(body, columns=headers), use_container_width=True)

st.divider()
c1, c2 = st.columns(2)
default_cur = c1.selectbox("통화 (파일에 통화 열이 없을 때)", ["KRW", "USD"])
default_dir = c2.selectbox("입/출금 구분이 비어 있을 때", ["출금", "입금"])

st.caption("💡 부서를 옮겨가며 열을 다 고치신 뒤, **마지막에 아래 버튼을 한 번만** 누르면 됩니다. "
           "고친 내용은 화면 안에 계속 유지됩니다.")
if st.button("📖 전부 읽어들이기", type="primary", use_container_width=True):
    allt = []
    for dept, _ in files:
        hr, mp = st.session_state[f"colmap_{dept}"]
        t = parse_plan_rows(info[dept]["rows"], mp, hr, period, dept, default_cur,
                            "out" if default_dir == "출금" else "in")
        save_column_mapping(dept, mp, mapping_signature(mp))
        allt.extend(t)
    st.session_state["plan_txns"] = allt
    st.success(f"{len(files)}개 부서에서 {len(allt)}건을 읽었습니다.")

txns = st.session_state.get("plan_txns")
if not txns:
    st.stop()

st.divider()
st.subheader("④ 읽어들인 내용")
bydept = pd.DataFrame([{"부서": d, "건수": sum(1 for t in txns if t.dept == d),
                        "입금": sum(t.amount_in for t in txns if t.dept == d),
                        "출금": sum(t.amount_out for t in txns if t.dept == d)}
                       for d in {t.dept for t in txns}]).sort_values("출금", ascending=False)
st.dataframe(bydept, use_container_width=True, hide_index=True)
st.caption("👆 부서별 원본 '합계' 열 총액과 맞는지 확인해주세요.")

dgroups = find_duplicates(txns)
if dgroups:
    st.subheader("⚠️ 같은 거래로 보이는 항목")
    for gi, group in enumerate(dgroups):
        st.write(f"**그룹 {gi+1}** — {txns[group[0]].memo} / {txns[group[0]].amount:,.0f}원")
        for i in group[1:]:
            if st.checkbox(f"{txns[i].dept} — {i+1}번째 건 제외", key=f"dup_{i}"):
                txns[i].excluded = 1

# ── 계정과목 확인 ─────────────────────────────────────────────────
st.divider()
st.subheader("⑤ 계정과목(태그) 확인")

with st.expander("❓ 계정과목은 어떻게 고르나요?"):
    st.markdown("""
부서 계획서에는 계정과목 칸이 없어서, **사유를 사전과 비교해 앱이 먼저 제안**합니다.

| 방법 | 이렇게 하세요 |
|---|---|
| 목록에서 고르기 | 오른쪽 칸을 눌러 펼친 뒤 클릭 |
| **키워드로 찾기** | 칸을 누른 채로 글자 입력 → 목록이 좁혀집니다 |

괄호 안 예시 단어로도 검색됩니다. 예) `통신` → `기타지출 (통신비·우편료·잡비)`
찾는 표현이 없으면 **'7 태그 사전 관리'** 에서 예시를 추가해두세요.

| 표시 | 뜻 | 할 일 |
|---|---|---|
| 🟢 자동 | 90점 이상 | 그냥 두셔도 됩니다 |
| 🟡 추천 | 70~89점 | 맞는지 확인 |
| 🔴 직접 선택 | 못 찾음 | 직접 골라주세요 |

한 번 고르면 사전에 쌓여, 다음 달부터 같은 사유는 자동 처리됩니다.
""")

_tags = load_tags()
tag_names = [""] + sorted(_tags.keys())
_hint = {"": "— 안 고름 —"}
for _n, _i in _tags.items():
    _ex = _i.example_list[:3]
    _hint[_n] = f"{_n}  ({'·'.join(_ex)})" if _ex else _n

PRIORITY = {"manual": 0, "suggest": 1, "auto": 2}
MARK = {"auto": "🟢", "suggest": "🟡", "manual": "🔴"}
live = [(i, t) for i, t in enumerate(txns) if not t.excluded]
sugs = {i: suggest_tag(t) for i, t in live}
order = sorted(live, key=lambda it: (PRIORITY[sugs[it[0]].level], -it[1].amount))
need = [it for it in order if sugs[it[0]].level != "auto"]

st.caption(f"확인 필요 {len(need)}건 · 자동 분류 {len(order)-len(need)}건")
show_all = st.checkbox("자동 분류된 것도 직접 고치기")
targets = order if show_all else need
if not targets:
    st.success("확인이 필요한 건이 없습니다. 아래 확인만 하고 저장하세요.")

PER_PAGE = 20
pages = max(1, (len(targets) + PER_PAGE - 1) // PER_PAGE)
page = st.number_input(f"페이지 (총 {pages}쪽 · 한 쪽 {PER_PAGE}건)", 1, pages, 1) if pages > 1 else 1
chunk = targets[(page - 1) * PER_PAGE: page * PER_PAGE]

for i, t in chunk:
    sug = sugs[i]
    c1, c2 = st.columns([3, 2])
    with c1:
        d = "입금" if t.amount_in else "출금"
        st.write(f"{MARK[sug.level]} **{t.memo}** — {t.counterpart} / {d} {t.amount:,.0f}원  \n"
                 f"<small>[{t.dept}] {sug.reason}</small>", unsafe_allow_html=True)
    with c2:
        idx = tag_names.index(sug.tag) if sug.tag in tag_names else 0
        t.tag = st.selectbox("계정과목", tag_names, index=idx,
                             format_func=lambda n: _hint.get(n, n),
                             key=f"tag_{i}", label_visibility="collapsed",
                             placeholder="입력해서 찾기 / 목록에서 고르기")

shown = {i for i, _ in chunk}
for i, t in live:
    if i not in shown and not t.tag:
        t.tag = sugs[i].tag

# ── 저장 전 확인 ──────────────────────────────────────────────────
st.divider()
st.subheader("⑥ 저장 전 확인 — 이렇게 분류됩니다")
st.caption("금액이 엉뚱한 계정과목에 몰려 있지 않은지 확인해주세요.")
summary = {}
for i, t in live:
    key = t.tag or "(안 고름 — 저장 안 됨)"
    e = summary.setdefault(key, {"건수": 0, "입금": 0.0, "출금": 0.0})
    e["건수"] += 1
    e["입금"] += t.amount_in
    e["출금"] += t.amount_out
st.dataframe(pd.DataFrame([{"계정과목": k, **v} for k, v in summary.items()])
             .sort_values("출금", ascending=False),
             use_container_width=True, hide_index=True)

with st.expander("전체 목록 보기", expanded=True):
    st.dataframe(pd.DataFrame([{
        "확인": MARK[sugs[i].level], "부서": t.dept,
        "계정과목": t.tag or "(안 고름)", "업체명": t.counterpart, "사유": t.memo,
        "통화": t.currency, "입금": t.amount_in, "출금": t.amount_out,
    } for i, t in live]), use_container_width=True, height=400)

no_tag = [t for _, t in live if not t.tag]
if no_tag:
    st.warning(f"계정과목을 안 고른 거래 {len(no_tag)}건은 저장되지 않습니다.")

# ── 이미 저장된 계획이 있는지 확인 ────────────────────────────────
existing = dept_counts(period, "plan")
mine = {d: existing.get(d, 0) for d in {t.dept for _, t in live} if existing.get(d, 0)}
if mine:
    st.warning("⚠️ **기존에 입력된 데이터가 있습니다. 갱신하시겠습니까?**  \n"
               + ", ".join(f"{d} {n}건" for d, n in mine.items())
               + "  \n저장 방식을 선택하지 않으면 :red[**중복 데이터**]가 발생할 수 있습니다.")

mode = st.radio(
    "저장 방식",
    ["갱신 — 기존 내용 :red[**삭제**] + 새 내용 저장  (권장)",
     "추가 — 기존 내용 :blue[**유지**] + 새 내용 덧붙임"],
    horizontal=True,
    help="같은 파일을 다시 올리거나 저장을 두 번 눌러도, '갱신'이면 중복이 생기지 않습니다.",
)

if st.button("💾 계획으로 저장하기", type="primary", use_container_width=True):
    keep = [t for _, t in live if t.tag]
    removed = 0
    if mode.startswith("갱신"):
        for d in {t.dept for t in keep}:
            removed += clear_period_dept(period, "plan", d)
    for t in keep:
        add_tag_example(t.tag, t.memo)
    n = save_transactions(keep)
    msg = f"계획 {n}건을 저장했습니다."
    if removed:
        msg += f" 기존 {removed}건은 이번 내용으로 바꿨습니다."
    if no_tag:
        msg += f" 계정과목을 안 고른 {len(no_tag)}건은 넣지 않았습니다."
    st.success(msg)
    st.session_state.pop("plan_txns", None)
