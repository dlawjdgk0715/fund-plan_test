"""자금실적 시트에 '마케팅,지급이자,외주용역비' 행을 한 번만 추가합니다 (방법 A).

원본 `2026년 자금실적` 시트의 C.지출 > 경상 아래, 'S/W, H/W 매입'(18행)과 '기타'(19행) 사이에
새 행을 끼워 넣고, 그 아래로 밀린 행을 가리키던 수식을 전부 한 칸씩 보정합니다.

사용법:
    python tools/add_marketing_row.py 원본파일.xlsm [출력파일.xlsm]

출력 파일을 templates/자금운용계획_template.xlsm 로 쓰시면 됩니다.
※ 한 번만 실행하세요. 두 번 실행하면 행이 두 개 생깁니다(실행 전에 확인합니다).
"""
from __future__ import annotations

import re
import sys
from copy import copy
from pathlib import Path

import openpyxl

SHEET = "2026년 자금실적"
INSERT_AT = 19                     # 이 행 자리에 새 행이 들어감 (기존 '기타'가 20으로 밀림)
NEW_LABEL = "마케팅,지급이자,외주용역비"
LABEL_COL = 4                      # D열 (인건비·S/W H/W 매입과 같은 단계)

REF = re.compile(r"(?<![A-Z0-9_$!'])(\$?)([A-Z]{1,3})(\$?)(\d+)")
SHEET_REF = re.compile(r"'([^']+)'!(\$?)([A-Z]{1,3})(\$?)(\d+)")


def shift_same_sheet(formula: str, at: int) -> str:
    """같은 시트를 가리키는 참조 중 at 이상인 행만 +1."""
    out, last = [], 0
    for m in SHEET_REF.finditer(formula):
        out.append(_shift(formula[last:m.start()], at))
        out.append(m.group(0))          # 다른 시트 참조는 그대로
        last = m.end()
    out.append(_shift(formula[last:], at))
    return "".join(out)


def _shift(text: str, at: int) -> str:
    def repl(m):
        row = int(m.group(4))
        return f"{m.group(1)}{m.group(2)}{m.group(3)}{row + 1 if row >= at else row}"
    return REF.sub(repl, text)


def shift_refs_to_sheet(formula: str, sheet: str, at: int) -> str:
    """특정 시트를 가리키는 참조 중 at 이상인 행만 +1 (다른 시트에서 참조할 때)."""
    def repl(m):
        if m.group(1) != sheet:
            return m.group(0)
        row = int(m.group(5))
        return f"'{sheet}'!{m.group(2)}{m.group(3)}{m.group(4)}{row + 1 if row >= at else row}"
    return SHEET_REF.sub(repl, formula)


def main(src: Path, dst: Path) -> None:
    wb = openpyxl.load_workbook(src, keep_vba=True)
    if SHEET not in wb.sheetnames:
        raise SystemExit(f"'{SHEET}' 시트를 못 찾았습니다.")
    ws = wb[SHEET]

    for r in range(1, ws.max_row + 1):
        if str(ws.cell(row=r, column=LABEL_COL).value or "").strip() == NEW_LABEL:
            raise SystemExit(f"이미 {r}행에 '{NEW_LABEL}' 행이 있습니다. 그대로 쓰시면 됩니다.")

    # 1) 대상 시트 안의 수식 먼저 보정
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                cell.value = shift_same_sheet(cell.value, INSERT_AT)

    # 2) 병합 셀을 잠시 풀고 행 삽입 (병합된 채로 밀면 글자가 사라집니다)
    from openpyxl.utils import get_column_letter as _L
    saved = [(r.min_row, r.min_col, r.max_row, r.max_col) for r in ws.merged_cells.ranges]
    for r in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(r))
    ws.insert_rows(INSERT_AT)
    for mn_r, mn_c, mx_r, mx_c in saved:
        sr = mn_r + 1 if mn_r >= INSERT_AT else mn_r
        er = mx_r + 1 if mx_r >= INSERT_AT else mx_r
        ws.merge_cells(f"{_L(mn_c)}{sr}:{_L(mx_c)}{er}")

    # 3) 새 행 채우기 — 서식은 바로 위 행(S/W, H/W 매입)에서 복사
    for col in range(1, ws.max_column + 1):
        above = ws.cell(row=INSERT_AT + 1, column=col)   # 밀려난 '기타' 행
        model = ws.cell(row=INSERT_AT - 1, column=col)   # S/W, H/W 매입 행
        new = ws.cell(row=INSERT_AT, column=col)
        if model.has_style:
            new.font = copy(model.font)
            new.border = copy(model.border)
            new.fill = copy(model.fill)
            new.number_format = model.number_format
            new.alignment = copy(model.alignment)
    ws.cell(row=INSERT_AT, column=LABEL_COL, value=NEW_LABEL)
    # 전년(2025년) 대응 행이 없으므로 합계·월평균은 0
    ws.cell(row=INSERT_AT, column=6, value=0)   # F열 합계
    ws.cell(row=INSERT_AT, column=7, value=0)   # G열 월평균

    # 4) 다른 시트에서 이 시트를 참조하던 수식 보정
    for other in wb.worksheets:
        if other.title == SHEET:
            continue
        for row in other.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and SHEET in cell.value:
                    cell.value = shift_refs_to_sheet(cell.value, SHEET, INSERT_AT)

    # 5) 원시 데이터 시트(6월데이터/7월데이터)의 수식 칸을 계산된 값으로 굳힘
    #    openpyxl로 저장하면 수식의 '마지막 계산값'이 지워져서, 엑셀에서 열기 전까지
    #    다른 프로그램이 그 칸을 빈 값으로 읽습니다. 이 두 시트는 매달 앱이 새로 덮어쓰는
    #    원시 데이터 자리라, 값으로 굳혀도 잃는 정보가 없습니다.
    wv = openpyxl.load_workbook(src, data_only=True)
    frozen = 0
    for name in ("6월데이터", "7월데이터"):
        if name not in wb.sheetnames or name not in wv.sheetnames:
            continue
        for row in wb[name].iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    cell.value = wv[name].cell(row=cell.row, column=cell.column).value
                    frozen += 1

    wb.save(dst)
    print(f"완료 — {INSERT_AT}행에 '{NEW_LABEL}' 행을 추가했습니다.")
    if frozen:
        print(f"      (원시 데이터 시트의 수식 칸 {frozen}개를 계산된 값으로 굳혔습니다)")
    print(f"저장: {dst}")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        raise SystemExit(__doc__)
    src = Path(sys.argv[1])
    dst = Path(sys.argv[2]) if len(sys.argv) > 2 else src.with_name(src.stem + "_행추가" + src.suffix)
    main(src, dst)
